"""Build an XLSX of employees missing their timesheet for one month — handed
to managers/ops so they can chase down who hasn't sent theirs, independent of
the Reminders page itself. Same visual conventions (header fill/font, frozen
header row, auto column width) as timesheet_export.py's build_timesheet_xlsx,
but a separate function: that one's shape is tightly coupled to the per-day
leave/status grid, which has nothing to do with this flat contact-list export.
"""
from __future__ import annotations

import calendar
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADERS: list[tuple[str, str]] = [
    ("employee_id", "Employee ID"),
    ("name", "Name"),
    ("account_manager", "Manager Name"),
    ("project", "Client"),
    ("personal_email", "Personal Email"),
    ("work_email", "Work Email"),
]


def build_missing_reminders_xlsx(rows: list[dict[str, Any]], month: int, year: int) -> bytes:
    """rows: one dict per employee missing this month's timesheet, keyed by
    _HEADERS' left-hand names."""
    wb = Workbook()
    ws = wb.active
    ws.title = f"Missing {calendar.month_abbr[month]} {year}"[:31]

    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(bold=True, color="FFFFFF")
    for col, (_, label) in enumerate(_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    for row_idx, data in enumerate(rows, 2):
        for col, (key, _) in enumerate(_HEADERS, 1):
            ws.cell(row=row_idx, column=col, value=data.get(key) or "")

    for col, (_, label) in enumerate(_HEADERS, 1):
        letter = get_column_letter(col)
        max_len = len(label)
        for row in range(2, len(rows) + 2):
            val = ws.cell(row=row, column=col).value
            if val:
                max_len = max(max_len, min(48, len(str(val))))
        ws.column_dimensions[letter].width = max(12, min(44, max_len + 2))

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
