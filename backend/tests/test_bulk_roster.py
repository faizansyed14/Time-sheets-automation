"""Bulk roster (one sheet, many employees) — code mapping, the per-row
checksum that makes the extraction trustworthy, and the deterministic
XLSX reader.

These are pure-unit: no DB, no LLM. The roster's own printed arithmetic
(Leave Days + Billing Days = Calendar Days) is what turns "the model read a
grid" into "the read is verified", so it gets the most coverage here.
"""
import io

import pytest

from app.services.extract_email import auto_accept
from app.services.bulk_roster.roster_codes import (
    distribute_days,
    leave_weight,
    normalise_code,
    verify_row_totals,
)
from app.services.bulk_roster.roster_parse import (
    looks_like_roster_grid,
    parse_period,
    parse_roster_cells,
)


def _roster_group(day_codes: dict[int, str], month: int = 7, year: int = 2026) -> dict:
    """The same {issues -> auto_accept} shape roster_stage.build_groups()
    actually produces for one row, without needing a DB session — enough to
    prove the row-level flags reach auto_accept.evaluate() and block it.

    Passes the row's OWN correct totals into verify_row_totals(), the way a
    real wide-format roster prints them (see test_no_stated_totals_... for
    the separate, deliberately-always-flagged case where a format prints no
    totals at all) — so these tests isolate the ambiguous-code/half-day
    blocking behavior specifically, without the "nothing was verified" flag
    also firing on every row here."""
    fields, uncertain, day_issues = distribute_days(day_codes, month, year)
    stated_leave = sum(leave_weight(c) for c in day_codes.values())
    issues = day_issues + verify_row_totals(day_codes, stated_leave, 31 - stated_leave, 31)
    return {
        "employee_pk": "pk-1", "name": "Test Employee", "employee_id": "E1",
        "month": month, "year": year,
        "sheets": [{"kind": "timesheet", "period_type": "full_month",
                    "days_covered": 31, "missing_days": [], "_days_in_month": 31}],
        "buckets": fields,
        "working_days": fields.get("working_days", []),
        "weekend_days": fields.get("weekend_days", []),
        "uncertain_days": uncertain, "issues": issues,
        "missing_days": [], "unaccounted_days": [],
        "days_covered_total": 31, "overlap_flags": [], "fold_notes": [],
    }


# --------------------------------------------------------------- code mapping
def test_codes_map_to_the_right_day_fields():
    fields, uncertain, issues = distribute_days(
        {1: "P", 2: "WK", 3: "L", 4: "AB", 5: "PH", 6: "HD", 7: "SL"}, 7, 2026)
    assert fields["working_days"] == ["2026-07-01"]
    assert fields["weekend_days"] == ["2026-07-02"]
    assert fields["annual"] == ["2026-07-03"]        # bare "L" -> annual, flagged elsewhere
    assert fields["absent"] == ["2026-07-04"]
    assert fields["public_holiday"] == ["2026-07-05"]
    assert fields["other"] == ["2026-07-06"]         # half day has no whole-day bucket
    assert fields["sick"] == ["2026-07-07"]
    # Days 8..31 were never supplied — they must surface as gaps, not vanish.
    assert len(uncertain) == 31 - 7
    assert any("not read" in i for i in issues)


def test_codes_are_case_and_punctuation_insensitive():
    assert normalise_code(" wk ") == "WK"
    assert normalise_code("P.") == "P"
    assert normalise_code("-") == ""
    assert normalise_code(None) == ""


def test_blank_and_unknown_cells_become_uncertain_never_a_guess():
    """A cell we can't read must reach a human, not get bucketed hopefully."""
    codes = {d: "P" for d in range(1, 32)}
    codes[10] = ""        # blank
    codes[11] = "ZZ"      # not a code we know
    fields, uncertain, issues = distribute_days(codes, 7, 2026)
    dates = {u["date"] for u in uncertain}
    assert dates == {"2026-07-10", "2026-07-11"}
    assert "2026-07-10" not in fields.get("working_days", [])
    assert "2026-07-11" not in fields.get("working_days", [])
    assert any("unrecognised" in i for i in issues)


def test_day_numbers_outside_the_month_are_dropped():
    """A fixed 1-31 form used for a 30-day month prints a dead column."""
    codes = {d: "P" for d in range(1, 31)}
    codes[31] = "P"
    fields, _uncertain, issues = distribute_days(codes, 6, 2026)   # June = 30 days
    assert "2026-06-31" not in fields["working_days"]
    assert any("outside" in i for i in issues)


def test_leave_weight_counts_half_days_as_half():
    assert leave_weight("L") == 1.0
    assert leave_weight("AB") == 1.0
    assert leave_weight("HD") == 0.5
    assert leave_weight("P") == 0.0
    assert leave_weight("WK") == 0.0
    assert leave_weight("PH") == 0.0


# ------------------------------------------------------------- the checksum
def test_matching_totals_produce_no_mismatch_issue():
    """Harshal Darak on the sample sheet: 5 leave days, 26 billing, 31 calendar."""
    codes = {d: "P" for d in range(1, 32)}
    for d in (13, 14, 15, 16, 17):
        codes[d] = "L"
    issues = verify_row_totals(codes, stated_leave=5, stated_billing=26, calendar_days=31)
    assert not any("mismatch" in i for i in issues)
    # A bare "L" is still surfaced as needing a leave-type confirmation.
    assert any('"L"' in i for i in issues)


def test_a_misread_cell_is_caught_by_the_rows_own_total():
    """THE point of this module: the grid says 4 leave days, the roster says
    5, so the row is flagged instead of quietly filed one day short."""
    codes = {d: "P" for d in range(1, 32)}
    for d in (13, 14, 15, 16):
        codes[d] = "L"
    issues = verify_row_totals(codes, stated_leave=5, stated_billing=26, calendar_days=31)
    assert any("leave-day mismatch" in i for i in issues)
    assert any("5" in i and "4" in i for i in issues)


def test_internally_inconsistent_roster_is_flagged():
    """leave + billing must equal calendar days; when the document's own two
    numbers disagree, no reading of it can be trusted either way."""
    codes = {d: "P" for d in range(1, 32)}
    issues = verify_row_totals(codes, stated_leave=0, stated_billing=25, calendar_days=31)
    assert any("don't add up" in i for i in issues)


def test_half_days_reconcile_and_are_still_flagged():
    codes = {d: "P" for d in range(1, 32)}
    codes[5] = "HD"
    codes[6] = "HD"
    issues = verify_row_totals(codes, stated_leave=1, stated_billing=30, calendar_days=31)
    assert not any("mismatch" in i for i in issues)   # 0.5 + 0.5 == 1
    assert any("half-day" in i for i in issues)


def test_no_stated_totals_means_no_false_mismatch_but_still_flags_for_review():
    """A roster without the summary columns simply can't be cross-checked —
    that must not manufacture a false MISMATCH claim, but it also must not
    look "clean" when nothing was actually verified: a format with no
    printed totals (e.g. PGC's "Certificate of Attendance") gets an explicit
    "review this directly" flag instead of silently qualifying for
    auto-accept on the strength of a check that never happened."""
    codes = {d: "P" for d in range(1, 32)}
    issues = verify_row_totals(codes, None, None, 31)
    assert not any("mismatch" in i or "don't add up" in i for i in issues)
    assert any("prints no Leave/Billing" in i for i in issues)


# ------------------------------------------- ambiguous L / HD forces review
def test_a_clean_row_with_no_ambiguity_may_auto_accept():
    codes = {d: "P" for d in range(1, 32)}
    codes[6] = codes[7] = "WK"
    codes[15] = "SL"     # sick has its own code — never ambiguous
    codes[20] = "AB"
    decision = auto_accept.evaluate(_roster_group(codes))
    assert decision.accepted is True, decision.blockers


def test_a_bare_l_code_blocks_auto_accept_for_manual_review():
    """A bare "L" doesn't say WHICH kind of leave — the roster's own notation
    is ambiguous, so this row must always go to a human, never auto-file."""
    codes = {d: "P" for d in range(1, 32)}
    codes[10] = "L"
    decision = auto_accept.evaluate(_roster_group(codes))
    assert decision.accepted is False
    assert any("marked only as" in b for b in decision.blockers)


def test_a_half_day_code_blocks_auto_accept_for_manual_review():
    """HD can't be represented as a whole leave day either way — always a
    human call, never auto-filed."""
    codes = {d: "P" for d in range(1, 32)}
    codes[10] = "HD"
    decision = auto_accept.evaluate(_roster_group(codes))
    assert decision.accepted is False
    assert any("half-day" in b for b in decision.blockers)


def test_l_or_hd_blocks_even_when_the_rows_own_totals_still_reconcile():
    """The gate must not be defeated just because the sheet's arithmetic
    happens to check out — an ambiguous code alone is reason enough."""
    codes = {d: "P" for d in range(1, 32)}
    codes[10] = codes[11] = "L"
    decision = auto_accept.evaluate(_roster_group(codes))
    # The checksum has nothing to disagree with (the row's own totals
    # reconcile perfectly), yet the row must still be blocked purely for the
    # L ambiguity.
    assert decision.accepted is False
    assert any("marked only as" in b for b in decision.blockers)


# ------------------------------------------------------- deterministic XLSX
def _sample_workbook() -> bytes:
    """A miniature of the real sheet: title row, legend, header, 3 people."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Approved Monthly Timesheet Jul 2026 (Contract Hire)"])
    ws.append(["Agency: Alphadata", "", "Calendar Days: 31"])
    ws.append([])
    header = ["Sr No", "Resource Name", "Title", "Loc.",
              "Emp. Timesheet Confirmation", "Leave Days", "Billing Days"]
    header += [f"{d:02d}" for d in range(1, 32)]
    ws.append(header)

    def person(sr, name, title, leave, billing, codes):
        row = [sr, name, title, "India", "Submitted", leave, billing]
        row += [codes.get(d, "P") for d in range(1, 32)]
        ws.append(row)

    person(1, "Abhishek Sharma", "HCM - Functional Consultant - I", 0, 31,
           {4: "WK", 5: "WK", 11: "WK", 12: "WK"})
    person(2, "Harshal Darak", "Finance - Technical Consultant - I", 5, 26,
           {13: "L", 14: "L", 15: "L", 16: "L", 17: "L", 4: "WK", 5: "WK"})
    person(3, "Vadde Murali", "Finance - Functional Consultant - II", 0, 31, {})

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_roster_is_read_without_any_llm():
    doc = parse_roster_cells("roster.xlsx", _sample_workbook())
    assert doc.method == "xlsx-cells"
    assert doc.llm_calls == 0
    assert (doc.month, doc.year) == (7, 2026)
    assert doc.calendar_days == 31
    assert doc.agency == "Alphadata"
    assert doc.headcount == 3
    assert [r.name for r in doc.rows] == [
        "Abhishek Sharma", "Harshal Darak", "Vadde Murali"]


def test_xlsx_row_carries_its_own_totals_and_full_day_grid():
    doc = parse_roster_cells("roster.xlsx", _sample_workbook())
    harshal = next(r for r in doc.rows if r.name == "Harshal Darak")
    assert harshal.sr_no == 2
    assert harshal.stated_leave_days == 5
    assert harshal.stated_billing_days == 26
    assert harshal.confirmation == "Submitted"
    assert harshal.location == "India"
    # Every one of the 31 day columns must be present — a partial read here
    # is the failure the whole flow is built to prevent.
    assert set(harshal.day_codes) == set(range(1, 32))
    assert [d for d, c in harshal.day_codes.items() if c == "L"] == [13, 14, 15, 16, 17]


def test_xlsx_read_reconciles_against_the_sheets_own_totals():
    """End-to-end on the deterministic path: what was read agrees with what
    the sheet printed, for every employee."""
    doc = parse_roster_cells("roster.xlsx", _sample_workbook())
    for r in doc.rows:
        issues = verify_row_totals(
            r.day_codes, r.stated_leave_days, r.stated_billing_days, doc.calendar_days)
        assert not any("mismatch" in i for i in issues), (r.name, issues)
        assert not any("don't add up" in i for i in issues), (r.name, issues)


def test_a_hundred_employee_roster_is_read_completely():
    """The headline scale question: 100 people on one sheet, nobody dropped
    and every row still reconciling against its own printed totals."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Approved Monthly Timesheet Jul 2026"])
    ws.append(["Calendar Days: 31"])
    header = ["Sr No", "Resource Name", "Title", "Loc.",
              "Emp. Timesheet Confirmation", "Leave Days", "Billing Days"]
    header += [f"{d:02d}" for d in range(1, 32)]
    ws.append(header)
    for i in range(1, 101):
        leave = i % 4                     # 0..3 leave days, varying per person
        codes = {d: "L" for d in range(10, 10 + leave)}
        row = [i, f"Employee {i:03d}", "Consultant", "India", "Submitted",
               leave, 31 - leave]
        row += [codes.get(d, "P") for d in range(1, 32)]
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)

    doc = parse_roster_cells("big.xlsx", buf.getvalue())
    assert doc.headcount == 100, "every employee row must be read, none skipped"
    assert [r.sr_no for r in doc.rows] == list(range(1, 101))
    for r in doc.rows:
        assert set(r.day_codes) == set(range(1, 32)), f"{r.name} lost day columns"
        issues = verify_row_totals(
            r.day_codes, r.stated_leave_days, r.stated_billing_days, doc.calendar_days)
        assert not any("mismatch" in i or "add up" in i for i in issues), (r.name, issues)


def test_duplicate_roster_rows_are_kept_once_and_reported():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Timesheet Jul 2026"])
    header = ["Sr No", "Resource Name", "Leave Days", "Billing Days"]
    header += [f"{d:02d}" for d in range(1, 32)]
    ws.append(header)
    for _ in range(2):
        ws.append([1, "Twice Listed", 0, 31] + ["P"] * 31)
    buf = io.BytesIO()
    wb.save(buf)

    doc = parse_roster_cells("dup.xlsx", buf.getvalue())
    assert doc.headcount == 1
    assert any("more than once" in i for i in doc.issues)


def test_a_non_roster_spreadsheet_is_rejected_not_mangled():
    """A single-employee timesheet must NOT be silently read as a roster —
    the caller falls back to the normal flow instead."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Employee Name", "Jane Doe"])
    ws.append(["Month", "July 2026"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError):
        parse_roster_cells("single.xlsx", buf.getvalue())


def test_period_is_parsed_from_several_title_shapes():
    assert parse_period("Approved Monthly Timesheet Jul 2026 (Contract Hire)") == (7, 2026)
    assert parse_period("Timesheet July-2026") == (7, 2026)
    assert parse_period("Monthly sheet 07/2026") == (7, 2026)
    assert parse_period("no period here") == (None, None)


# ---------------------------------------------- long format ("PGC"-style)
def _long_format_workbook(rows: list[tuple]) -> bytes:
    """rows: (employee_id, name, title, date, attendance_type) tuples — a
    miniature of a real long/flat export (one row per employee PER DAY)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["EmployeeId", "EmployeeName", "JobTitle", "AttendanceDate",
               "DayName", "AttendanceType"])
    for emp_id, name, title, date, atype in rows:
        ws.append([emp_id, name, title, date, "", atype])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_long_format_is_detected_and_read_without_any_llm():
    import datetime as dt

    data = _long_format_workbook([
        ("E1", "Ada Lovelace", "Engineer", dt.date(2026, 7, 1), "Daily Present"),
        ("E1", "Ada Lovelace", "Engineer", dt.date(2026, 7, 2), "Weekend"),
        ("E2", "Grace Hopper", "Engineer", "03-Jul-2026", "Sick Leave"),
    ])
    doc = parse_roster_cells("pgc.xlsx", data)
    assert doc.method == "xlsx-cells-long"
    assert doc.llm_calls == 0
    assert (doc.month, doc.year) == (7, 2026)
    assert doc.headcount == 2


def test_long_format_does_not_misfire_on_the_wide_format_sample():
    """The two shapes must never be confused for each other — detection is
    checked against BOTH false directions, not just "does PGC parse"."""
    doc = parse_roster_cells("roster.xlsx", _sample_workbook())
    assert doc.method == "xlsx-cells"


def test_wide_format_with_a_status_column_is_still_not_misdetected():
    """A wide sheet whose confirmation column happens to be literally named
    "Status" (a synonym shared with the long format's attendance-type
    pattern) plus an "Employee ID" column must still resolve as wide — there
    is no "date" column on it, so the four-way AND can never be satisfied."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Timesheet Jul 2026"])
    header = ["Sr No", "Employee Name", "Employee ID", "Status", "Leave Days", "Billing Days"]
    header += [f"{d:02d}" for d in range(1, 32)]
    ws.append(header)
    ws.append([1, "Jane Doe", "E100", "Submitted", 0, 31] + ["P"] * 31)
    buf = io.BytesIO()
    wb.save(buf)
    doc = parse_roster_cells("wide_status.xlsx", buf.getvalue())
    assert doc.method == "xlsx-cells"


def test_long_format_maps_the_four_known_attendance_types():
    import datetime as dt

    data = _long_format_workbook([
        ("E1", "Ada Lovelace", "Engineer", dt.date(2026, 7, 1), "Daily Present"),
        ("E1", "Ada Lovelace", "Engineer", dt.date(2026, 7, 2), "Weekend"),
        ("E1", "Ada Lovelace", "Engineer", dt.date(2026, 7, 3), "Sick Leave"),
        ("E1", "Ada Lovelace", "Engineer", dt.date(2026, 7, 4), "Annual Leave"),
    ])
    doc = parse_roster_cells("pgc.xlsx", data)
    row = doc.rows[0]
    fields, _uncertain, _issues = distribute_days(row.day_codes, doc.month, doc.year)
    assert fields["working_days"] == ["2026-07-01"]
    assert fields["weekend_days"] == ["2026-07-02"]
    assert fields["sick"] == ["2026-07-03"]
    assert fields["annual"] == ["2026-07-04"]


def test_long_format_unfamiliar_type_becomes_uncertain_never_guessed():
    """"Internal Mission (Inside UAE)" etc. are NOT in the known table on
    purpose — they must reach a human, never get silently bucketed."""
    import datetime as dt

    data = _long_format_workbook([
        ("E1", "Ada Lovelace", "Engineer", dt.date(2026, 7, 1), "Internal Mission (Inside UAE)"),
    ])
    doc = parse_roster_cells("pgc.xlsx", data)
    row = doc.rows[0]
    assert row.day_codes[1] == "Internal Mission (Inside UAE)"   # kept raw, not guessed
    _fields, uncertain, issues = distribute_days(row.day_codes, doc.month, doc.year)
    assert "2026-07-01" in {u["date"] for u in uncertain}
    assert any("Internal Mission" in i for i in issues)


def test_long_format_pending_approval_suffix_is_bucketed_and_flagged():
    import datetime as dt

    data = _long_format_workbook([
        ("E1", "Ada Lovelace", "Engineer", dt.date(2026, 7, 1),
         "Sick Leave - Waiting for Approval -   Line Manager"),
    ])
    doc = parse_roster_cells("pgc.xlsx", data)
    row = doc.rows[0]
    assert row.day_codes[1] == "SL"                              # still classified as sick
    assert any("not yet approved" in i for i in row.issues)      # but flagged for confirmation


def test_long_format_real_dates_and_string_dates_both_parse():
    import datetime as dt

    data = _long_format_workbook([
        ("E1", "Ada Lovelace", "Engineer", dt.date(2026, 7, 15), "Daily Present"),
        ("E2", "Grace Hopper", "Engineer", "15-Jul-2026", "Daily Present"),
    ])
    doc = parse_roster_cells("pgc.xlsx", data)
    for row in doc.rows:
        assert row.day_codes.get(15) == "P"


def test_long_format_groups_by_employee_not_raw_row_count():
    """Hundreds of raw rows for a handful of employees must collapse to one
    RosterRow per employee — this is the scale case the format exists for."""
    import datetime as dt

    rows = []
    for emp_n in range(1, 6):          # 5 employees
        for day in range(1, 29):       # 28 days each = 140 raw rows
            rows.append((f"E{emp_n}", f"Employee {emp_n}", "Engineer",
                         dt.date(2026, 7, day), "Daily Present"))
    data = _long_format_workbook(rows)
    doc = parse_roster_cells("pgc.xlsx", data)
    assert len(rows) == 140
    assert doc.headcount == 5
    for row in doc.rows:
        assert len(row.day_codes) == 28


def test_long_format_bad_or_duplicate_dates_are_flagged_not_dropped():
    """A row with no readable date, a duplicated day, and a day from a
    different month must each surface as an issue on that employee — the
    employee themselves is still staged, never silently skipped."""
    import datetime as dt

    data = _long_format_workbook([
        ("E1", "Test Person", "Officer", dt.date(2026, 7, 1), "Daily Present"),
        ("E1", "Test Person", "Officer", None, "Weekend"),                     # unreadable date
        ("E1", "Test Person", "Officer", dt.date(2026, 7, 1), "Weekend"),      # duplicate day 1
        ("E1", "Test Person", "Officer", dt.date(2026, 8, 15), "Weekend"),     # wrong month
    ])
    doc = parse_roster_cells("pgc.xlsx", data)
    assert doc.headcount == 1
    row = doc.rows[0]
    assert row.day_codes == {1: "P"}
    assert any("no readable AttendanceDate" in i for i in row.issues)
    assert any("more than once" in i for i in row.issues)
    assert any("doesn't belong to" in i for i in row.issues)


def test_long_format_translate_attendance_type_directly():
    from app.services.bulk_roster.roster_codes import translate_attendance_type

    assert translate_attendance_type("Daily Present") == ("P", None)
    assert translate_attendance_type("Weekend") == ("WK", None)
    assert translate_attendance_type("Sick Leave") == ("SL", None)
    assert translate_attendance_type("Annual Leave - Completed") == ("AL", None)

    code, flag = translate_attendance_type("Sick Leave - Waiting for Approval -   Line Manager")
    assert code == "SL"
    assert flag and "not yet approved" in flag

    code, flag = translate_attendance_type("Internal Mission (Inside UAE)")
    assert code == "Internal Mission (Inside UAE)"
    assert flag is None

    assert translate_attendance_type(None) == ("", None)
    assert translate_attendance_type("") == ("", None)


# ------------------------------------- the /upload page's roster guard
def test_roster_guard_catches_both_shapes_but_not_a_real_single_employee_sheet():
    """The deterministic, zero-LLM guard used to keep a mistakenly-uploaded
    roster off the single-employee /upload page (see api/routes/upload.py) —
    must catch either sheet shape, and must NOT flag a genuine one-person
    timesheet, which is the whole point of the pipeline it protects."""
    import datetime as dt

    long_data = _long_format_workbook([
        ("E1", "Ada Lovelace", "Engineer", dt.date(2026, 7, 1), "Daily Present"),
    ])
    assert looks_like_roster_grid("pgc.xlsx", long_data) is True
    assert looks_like_roster_grid("roster.xlsx", _sample_workbook()) is True

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Employee Name", "Jane Doe"])
    ws.append(["Month", "July 2026"])
    buf = io.BytesIO()
    wb.save(buf)
    assert looks_like_roster_grid("single.xlsx", buf.getvalue()) is False


def test_roster_guard_never_crashes_on_a_non_spreadsheet_file():
    """A file that isn't a spreadsheet at all is a different problem — the
    guard must defer to the normal pipeline, not raise."""
    assert looks_like_roster_grid("notes.pdf", b"not a spreadsheet at all") is False
    assert looks_like_roster_grid("photo.jpg", b"\xff\xd8\xff\xe0") is False


# --------------------------------------------------------------------------- #
# FAZAA-style "Organization-Wise Attendance" reports — a PDF, one row per
# employee PER DAY, punch times + an unreliable AB/PR/WO notation + a
# "Reason" annotation, read via text (not vision/image — see
# roster_extract._vision_attendance_report's own docstring for why) and
# classified in code (roster_codes.translate_attendance_report_row).
# --------------------------------------------------------------------------- #
from app.services.bulk_roster.roster_codes import translate_attendance_report_row
from app.services.bulk_roster.roster_extract import (
    _attendance_report_page_texts,
    _looks_like_attendance_report,
)


def _fitz_pdf(lines: list[str]) -> bytes:
    """A minimal one-page PDF with the given lines of text — just enough for
    fitz's own text extraction to read something back, not a realistic
    render of the actual report."""
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    y = 50
    for line in lines:
        page.insert_text((50, y), line)
        y += 20
    buf = doc.tobytes()
    doc.close()
    return buf


def test_attendance_report_fingerprint_fires_only_on_the_real_combination():
    real = _fitz_pdf([
        "FAZAA", "Organization-Wise Attendance From 01/07/2026 To 31/07/2026",
        "User ID Name Shift IN-SPFID OUT-SPFID 1st Half 2nd Half",
    ])
    assert _looks_like_attendance_report(real) is True

    # Title alone, no SPFID vocabulary — must NOT fire (a different roster
    # style could plausibly share generic attendance-report wording).
    title_only = _fitz_pdf(["Organization-Wise Attendance From 01/07/2026 To 31/07/2026"])
    assert _looks_like_attendance_report(title_only) is False

    # SPFID alone, no title — must NOT fire either.
    spfid_only = _fitz_pdf(["User ID Name Shift IN-SPFID OUT-SPFID"])
    assert _looks_like_attendance_report(spfid_only) is False

    # An unrelated roster/PDF, and outright non-PDF bytes — never raises.
    unrelated = _fitz_pdf(["Approved Monthly Timesheet - July 2026", "Sr No Name Title"])
    assert _looks_like_attendance_report(unrelated) is False
    assert _looks_like_attendance_report(b"not a pdf at all") is False
    assert _looks_like_attendance_report(b"") is False


def test_attendance_report_page_texts_reconstructs_rows_by_position():
    """Word-position reconstruction — not asserting exact spacing, just that
    words on the same visual line end up on the same output line, in
    left-to-right order, which is the property the reader's prompt depends
    on for a Reason annotation to land next to its own row."""
    pdf = _fitz_pdf(["FAC13763 07:45 15:30 AB AB"])
    texts = _attendance_report_page_texts(pdf)
    assert len(texts) == 1
    words = texts[0].split()
    assert words == ["FAC13763", "07:45", "15:30", "AB", "AB"]


def test_attendance_report_page_texts_handles_a_blank_page():
    import fitz

    doc = fitz.open()
    doc.new_page()  # no text at all
    buf = doc.tobytes()
    doc.close()
    assert _attendance_report_page_texts(buf) == [""]


# ---- translate_attendance_report_row: the actual bucket decision ----

def test_translate_attendance_report_row_present_when_real_times_punched():
    assert translate_attendance_report_row("07:45", "15:30", "AB", "AB", None) == "P"


def test_translate_attendance_report_row_ignores_ab_pr_as_a_presence_signal():
    """Confirmed against real client data: AB/PR do not reliably mean
    absent/present on this report — real punch times must win regardless."""
    assert translate_attendance_report_row("07:45", "15:30", "AB", "AB", None) == "P"
    assert translate_attendance_report_row(None, None, "PR", "PR", None) == ""


def test_translate_attendance_report_row_weekly_off():
    assert translate_attendance_report_row(None, None, "WO", "WO", None) == "WO"
    # WO can land in the wrong field slot when times are blank — still caught.
    assert translate_attendance_report_row("WO", "WO", None, None, None) == "WO"


def test_translate_attendance_report_row_known_reasons_map_to_existing_codes():
    assert translate_attendance_report_row(None, None, None, None, "Sick leave") == "SL"
    assert translate_attendance_report_row(None, None, None, None, "Annual Leave") == "AL"
    # Case/whitespace-insensitive.
    assert translate_attendance_report_row(None, None, None, None, "  SICK LEAVE  ") == "SL"


def test_translate_attendance_report_row_unmapped_reasons_pass_through_uncertain():
    """"Time off - Personal"/"Time off - Business"/"Event Leave"/"Forget to
    stamp out" have no existing bucket — never guessed, passed through
    verbatim so distribute_days() turns them into an uncertain_days entry
    with the reviewer seeing the exact original text."""
    assert translate_attendance_report_row(None, None, None, None, "Time off personal") == "Time off personal"
    assert translate_attendance_report_row(None, None, None, None, "Forget to stamp out") == "Forget to stamp out"


def test_translate_attendance_report_row_blank_row_with_no_signal_is_blank():
    assert translate_attendance_report_row(None, None, None, None, None) == ""


def test_translate_attendance_report_row_never_trusts_garbled_non_time_fields_as_present():
    """The real, confirmed bug this guard exists for: the reader can
    misattribute free text (e.g. "Sick leave") into the in_time/out_time
    slots when a row's real times are blank. Without validating the SHAPE of
    those fields, a non-null-but-garbled in_time/out_time would wrongly
    resolve to "P" — silently filing a full-day absence as a normal working
    day. Must land blank (uncertain) instead."""
    assert translate_attendance_report_row("Sick leave", "AB", None, None, None) == ""
    assert translate_attendance_report_row("14:5", "15:30", None, None, None) != "P"  # malformed time shape


# ---- end-to-end: detection -> two independent reads -> merge/corroborate ----

async def test_attendance_report_end_to_end_corroborates_reads_and_flags_disagreement(mock_vision_calls):
    from app.services.bulk_roster.roster_extract import extract_roster

    pdf = _fitz_pdf([
        "FAZAA", "Organization-Wise Attendance From 01/07/2026 To 31/07/2026",
        "User ID Name Shift IN-SPFID OUT-SPFID 1st Half 2nd Half",
        "01/07/2026 (Wednesday)",
        "E1 Ada Lovelace GS 07:45 15:30 AB AB",
        "02/07/2026 (Thursday)",
        "E1 Ada Lovelace GS AB AB",
    ])

    # Both reads AGREE on 07-01 (a real working day) — trusted as-is.
    # Both reads DISAGREE on 07-02 — one says Sick Leave, the other blank —
    # must never be trusted either way.
    read1 = {"rows": [
        {"employee_id": "E1", "name": "Ada Lovelace", "date": "2026-07-01",
         "in_time": "07:45", "out_time": "15:30", "half1": "AB", "half2": "AB", "reason": None},
        {"employee_id": "E1", "name": "Ada Lovelace", "date": "2026-07-02",
         "in_time": None, "out_time": None, "half1": "AB", "half2": "AB", "reason": "Sick leave"},
    ], "rows_visible": 2}
    read2 = {"rows": [
        {"employee_id": "E1", "name": "Ada Lovelace", "date": "2026-07-01",
         "in_time": "07:45", "out_time": "15:30", "half1": "AB", "half2": "AB", "reason": None},
        {"employee_id": "E1", "name": "Ada Lovelace", "date": "2026-07-02",
         "in_time": None, "out_time": None, "half1": "AB", "half2": "AB", "reason": None},
    ], "rows_visible": 2}
    mock_vision_calls([read1, read2])

    doc = await extract_roster("rptDtAttendance.pdf", pdf)

    assert doc.method == "vision-attendance-report"
    assert doc.month == 7 and doc.year == 2026
    assert len(doc.rows) == 1
    row = doc.rows[0]
    assert row.employee_id == "E1"
    assert row.day_codes[1] == "P"     # agreed -> trusted
    assert row.day_codes[2] == ""      # disagreed (SL vs blank) -> never trusted, uncertain
    assert any("disagreed" in issue for issue in row.issues)


# --------------------------------------------------------------------------- #
# HHRC-style "Outsourced Staff Attendance Report" — a genuinely clean XLSX
# (real typed datetime/time cells), long-format (one row per employee PER
# DAY) with a Type + Sub Type pair instead of a single code. Zero LLM calls,
# same class as the PGC long-format reader above — see roster_parse.
# _parse_hhrc_format / roster_codes.translate_hhrc_day for the design.
# --------------------------------------------------------------------------- #
from app.services.bulk_roster.roster_codes import translate_hhrc_day
from app.services.bulk_roster.roster_parse import _detect_hhrc_format, _load_grid


def _hhrc_workbook(rows: list[tuple]) -> bytes:
    """rows: (employee_number, name, date, type, subtype, time_in, time_out) —
    a miniature of the real HHRC export's own column layout."""
    import datetime as dt

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Employee Number", "Employee Name", "Employee Email", "Organization",
               "Date In", "Date Out", "Type", "Sub Type", "Time In", "Time Out",
               "Work Remotely Flag", "Hours", "Job Title", "Shift Name",
               "Supervisor Number", "Supervisor Name"])
    for emp_id, name, date, typ, subtype, tin, tout in rows:
        ws.append([emp_id, name, f"{name}@example.ae", "IT Department",
                   date, date, typ, subtype, tin, tout, "No",
                   dt.time(0, 0), "IT Consultant", "Normal Shift", 999, "Some Manager"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_hhrc_format_is_detected_and_read_without_any_llm():
    import datetime as dt

    data = _hhrc_workbook([
        (500003, "Salman Ahmed", dt.datetime(2026, 7, 1), "Normal", None,
         dt.time(7, 45), dt.time(15, 30)),
        (500003, "Salman Ahmed", dt.datetime(2026, 7, 2), "Week End", None,
         dt.time(0, 0), dt.time(0, 0)),
        (500003, "Salman Ahmed", dt.datetime(2026, 7, 3), "On Leave", "ANNUAL LEAVE",
         dt.time(0, 0), dt.time(0, 0)),
    ])
    doc = parse_roster_cells("hhrc.xlsx", data)
    assert doc.method == "xlsx-cells-hhrc"
    assert doc.llm_calls == 0
    assert (doc.month, doc.year) == (7, 2026)
    assert doc.headcount == 1
    row = doc.rows[0]
    assert row.employee_id == "500003"
    assert row.sr_no == 500003
    assert row.day_codes == {1: "P", 2: "WK", 3: "AL"}


def test_hhrc_format_never_misfires_on_pgc_long_format_or_the_wide_format():
    """Detection is checked against every OTHER shape, not just "does HHRC
    parse" — the three formats' header vocabularies must never collide."""
    import datetime as dt

    pgc_data = _long_format_workbook([
        ("E1", "Ada Lovelace", "Engineer", dt.date(2026, 7, 1), "Daily Present"),
    ])
    assert _detect_hhrc_format(_load_grid("pgc.xlsx", pgc_data)) is None

    assert _detect_hhrc_format(_load_grid("roster.xlsx", _sample_workbook())) is None

    hhrc_data = _hhrc_workbook([
        (500003, "Salman Ahmed", dt.datetime(2026, 7, 1), "Normal", None,
         dt.time(7, 45), dt.time(15, 30)),
    ])
    assert _detect_hhrc_format(_load_grid("hhrc.xlsx", hhrc_data)) is not None
    from app.services.bulk_roster.roster_parse import _detect_long_format
    assert _detect_long_format(_load_grid("hhrc.xlsx", hhrc_data)) is None


def test_hhrc_format_skips_a_trailing_footer_row_with_no_name():
    """The real file this reader was built against has a trailing
    "Information Classification: Confidential" row (text in the Employee
    Number column, every other column blank) — must be silently skipped via
    the same blank-name guard the PGC reader already uses, not treated as a
    14th employee."""
    import datetime as dt

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Employee Number", "Employee Name", "Employee Email", "Organization",
               "Date In", "Date Out", "Type", "Sub Type", "Time In", "Time Out",
               "Work Remotely Flag", "Hours", "Job Title", "Shift Name",
               "Supervisor Number", "Supervisor Name"])
    ws.append([500003, "Salman Ahmed", "s@example.ae", "IT", dt.datetime(2026, 7, 1),
               dt.datetime(2026, 7, 1), "Normal", None, dt.time(7, 45), dt.time(15, 30),
               "No", dt.time(7, 45), "IT Consultant", "Shift", 999, "Manager"])
    ws.append(["Information Classification: Confidential", None, None, None,
               None, None, None, None, None, None, None, None, None, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)

    doc = parse_roster_cells("hhrc.xlsx", buf.getvalue())
    assert doc.headcount == 1
    assert doc.rows[0].employee_id == "500003"


def test_hhrc_format_merges_multiple_rows_for_the_same_employee_day():
    """The one real wrinkle this format has: a short Personal/Official
    segment recorded as its OWN row alongside a Normal row for the same
    calendar day — confirmed real cases in the actual file. Both rows must
    merge into ONE day_codes entry, and a real Normal segment must win the
    day's bucket regardless of what else shares that day."""
    import datetime as dt

    data = _hhrc_workbook([
        (500024, "Rehab Ghoneim", dt.datetime(2026, 7, 13), "Normal", None,
         dt.time(8, 32), dt.time(16, 0)),
        (500024, "Rehab Ghoneim", dt.datetime(2026, 7, 13), "Personal", None,
         dt.time(8, 30), dt.time(8, 32)),
        (500024, "Rehab Ghoneim", dt.datetime(2026, 7, 28), "Official", None,
         dt.time(8, 30), dt.time(10, 13)),
        (500024, "Rehab Ghoneim", dt.datetime(2026, 7, 28), "Normal", None,
         dt.time(10, 13), dt.time(15, 36)),
    ])
    doc = parse_roster_cells("hhrc.xlsx", data)
    row = doc.rows[0]
    assert row.day_codes[13] == "P"
    assert row.day_codes[28] == "P"


# ---- translate_hhrc_day: the actual per-day bucket decision ----

def test_translate_hhrc_day_normal_alone():
    assert translate_hhrc_day([("Normal", None)]) == "P"


def test_translate_hhrc_day_normal_wins_over_any_other_segment_sharing_the_day():
    assert translate_hhrc_day([("Personal", None), ("Normal", None)]) == "P"
    assert translate_hhrc_day([("Official", None), ("Normal", None)]) == "P"


def test_translate_hhrc_day_week_end():
    assert translate_hhrc_day([("Week End", None)]) == "WK"


def test_translate_hhrc_day_on_leave_known_subtype():
    assert translate_hhrc_day([("On Leave", "ANNUAL LEAVE")]) == "AL"
    # Case/whitespace-insensitive.
    assert translate_hhrc_day([("On Leave", "  annual leave  ")]) == "AL"


def test_translate_hhrc_day_on_leave_unknown_subtype_is_uncertain_not_guessed():
    """No dedicated bucket for an unrecognised Sub Type — surfaced verbatim
    so distribute_days() turns it into an uncertain_days entry, same "never
    guess" treatment as every other format's unmapped reason/code."""
    assert translate_hhrc_day([("On Leave", "SICK LEAVE")]) == "SICK LEAVE"
    assert translate_hhrc_day([("On Leave", None)]) == "On Leave"


def test_translate_hhrc_day_personal_alone_with_no_normal_is_never_guessed():
    """Never observed in the real file (Personal/Official always come with
    a Normal segment) — must not be silently assumed to mean "present"."""
    result = translate_hhrc_day([("Personal", None)])
    assert result not in ("P", "WK", "AL")
