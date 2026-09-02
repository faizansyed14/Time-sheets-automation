"""Generic AI chat — attachment extraction + message building (deterministic,
no LLM) + endpoints.

The chat itself has no database access and no tools any more — it's a plain
general-purpose assistant that can read an attached file. Nothing here needs
a live model: extraction and message-building are pure functions, and the
endpoint tests run with no AI provider configured (the standard test env),
which exercises the graceful-degrade path without needing a real API key.
"""
import io
import json

import pytest
from fastapi import UploadFile

from app.api.routes.agentic_chat import _read_capped
from app.services.agents import chat_agent, chat_files
from tests.conftest import auth_headers


# --------------------------------------------------------------------- #
# agentic_chat._read_capped — pure, no DB/LLM
# --------------------------------------------------------------------- #
async def test_read_capped_returns_none_for_a_file_over_the_cap():
    f = UploadFile(io.BytesIO(b"x" * 100))
    assert await _read_capped(f, cap=50) is None


async def test_read_capped_returns_bytes_for_a_file_within_the_cap():
    f = UploadFile(io.BytesIO(b"hello world"))
    assert await _read_capped(f, cap=1_000_000) == b"hello world"


# --------------------------------------------------------------------- #
# chat_files.extract_attachment — pure, no DB/LLM
# --------------------------------------------------------------------- #
def test_extract_plain_text_file():
    result = chat_files.extract_attachment("notes.txt", "text/plain", b"hello world\nline two")
    assert result["kind"] == "text"
    assert "hello world" in result["text"]
    assert result["truncated"] is False


def test_extract_pdf_text():
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello from a PDF")
    data = doc.tobytes()
    doc.close()
    result = chat_files.extract_attachment("sample.pdf", "application/pdf", data)
    assert result["kind"] == "text"
    # Words on the same line are tab-joined (not space-joined) by design — see
    # test_extract_pdf_preserves_table_row_alignment for why — so check the
    # words appear, in order, rather than an exact space-joined phrase.
    assert all(w in result["text"] for w in ("Hello", "from", "a", "PDF"))
    assert result["text"].index("Hello") < result["text"].index("PDF")


def test_extract_pdf_preserves_table_row_alignment():
    """THE actual bug this replaces: PyMuPDF's default page.get_text() flattens
    a dense day-by-day table into one value per line with no row grouping,
    forcing a model to count position across dozens of disconnected lines to
    figure out which day a value belongs to — that's exactly how a real
    off-by-one leave-date misread happened. Row-by-row reconstruction from
    each word's own coordinates must keep the header row and a data row each
    on their own line, with matching column position."""
    import fitz
    doc = fitz.open()
    page = doc.new_page(width=1000, height=200)
    headers = ["Sr", "Name"] + [f"{d:02d}" for d in range(1, 8)]
    x = 20
    for h in headers:
        page.insert_text((x, 30), h, fontsize=6)
        x += 22
    row = ["4", "Person"] + ["P", "P", "L", "L", "L", "WK", "WK"]
    x = 20
    for v in row:
        page.insert_text((x, 45), v, fontsize=6)
        x += 22
    data = doc.tobytes()
    doc.close()

    result = chat_files.extract_attachment("roster.pdf", "application/pdf", data)
    lines = [ln for ln in result["text"].splitlines() if ln.strip()]
    assert len(lines) == 2, f"expected header + one data row, got: {lines!r}"
    header_cols = lines[0].split("\t")
    data_cols = lines[1].split("\t")
    # "03" (day 3) and "04" (day 4) in the header must line up with the "L"s
    # in the data row at the SAME column position — not shifted by one.
    assert header_cols[header_cols.index("03")] == "03"
    assert data_cols[header_cols.index("03")] == "L"
    assert data_cols[header_cols.index("04")] == "L"
    assert data_cols[header_cols.index("02")] == "P"
    assert data_cols[header_cols.index("06")] == "WK"


def test_extract_scanned_pdf_errors_clearly_no_vision_fallback():
    """A scan (a photo of a document, no embedded text layer) has nothing
    for _extract_pdf_text to find. This module deliberately does NOT fall
    back to vision for that case — a clear "no readable text" error is the
    whole contract here, matching every other unreadable-file case."""
    pytest.importorskip("PIL")
    from PIL import Image
    img_buf = io.BytesIO()
    Image.new("RGB", (200, 100), color="white").save(img_buf, format="PNG")

    import fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_image(fitz.Rect(0, 0, 200, 100), stream=img_buf.getvalue())
    data = doc.tobytes()
    doc.close()

    result = chat_files.extract_attachment("scan.pdf", "application/pdf", data)
    assert result["kind"] == "error"
    assert "no readable text" in result["message"].lower()


def test_extract_docx_text():
    import docx
    d = docx.Document()
    d.add_paragraph("Hello from a DOCX")
    d.add_table(rows=1, cols=2).rows[0].cells[0].text = "cell one"
    buf = io.BytesIO()
    d.save(buf)
    result = chat_files.extract_attachment(
        "sample.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        buf.getvalue(),
    )
    assert result["kind"] == "text"
    assert "Hello from a DOCX" in result["text"]
    assert "cell one" in result["text"]


def test_extract_xlsx_text():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Hours"])
    ws.append(["Faizan", 40])
    buf = io.BytesIO()
    wb.save(buf)
    result = chat_files.extract_attachment(
        "sample.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        buf.getvalue(),
    )
    assert result["kind"] == "text"
    assert "Faizan" in result["text"] and "40" in result["text"]


def test_extract_image_is_sent_as_image_not_extracted_as_text():
    pytest.importorskip("PIL")
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (2, 2), color="red").save(buf, format="PNG")
    result = chat_files.extract_attachment("pic.png", "image/png", buf.getvalue())
    assert result["kind"] == "image"
    assert result["mime"] == "image/png"
    assert result["b64"]


def test_extract_unsupported_legacy_formats_error_clearly_instead_of_garbage_text():
    result = chat_files.extract_attachment("old.doc", "application/msword", b"not a real doc")
    assert result["kind"] == "error"
    assert ".docx" in result["message"]

    result2 = chat_files.extract_attachment("old.xls", "application/vnd.ms-excel", b"not a real xls")
    assert result2["kind"] == "error"
    assert ".xlsx" in result2["message"]


def test_extract_empty_file_errors_rather_than_silently_passing_nothing():
    result = chat_files.extract_attachment("empty.txt", "text/plain", b"   \n  ")
    assert result["kind"] == "error"


def test_extract_rejects_bytes_that_are_not_actually_a_valid_image():
    """An attacker-controlled extension/content-type proves nothing on its
    own — garbage bytes labeled .png must be rejected, never silently
    base64'd and handed to the model as a real image."""
    result = chat_files.extract_attachment("fake.png", "image/png", b"not an image, just plain bytes")
    assert result["kind"] == "error"
    assert "valid image" in result["message"]


def test_extract_accepts_a_genuinely_valid_image():
    pytest.importorskip("PIL")
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (2, 2), color="blue").save(buf, format="PNG")
    result = chat_files.extract_attachment("pic.png", "image/png", buf.getvalue())
    assert result["kind"] == "image"


def test_zip_bomb_guard_rejects_a_declared_size_over_the_cap(monkeypatch):
    """A normal, tiny XLSX must still pass; lowering the cap below its own
    declared size proves the size-comparison itself actually rejects,
    without needing to construct a real 200MB zip bomb to prove it."""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["a", "b", "c"])
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    # Passes at the real cap.
    ok = chat_files.extract_attachment(
        "normal.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", data)
    assert ok["kind"] == "text"

    # Same file, cap lowered below its own declared uncompressed size -> rejected.
    monkeypatch.setattr(chat_files, "MAX_ZIP_UNCOMPRESSED_BYTES", 1)
    blocked = chat_files.extract_attachment(
        "normal.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", data)
    assert blocked["kind"] == "error"
    assert "too large" in blocked["message"]


def test_zip_bomb_guard_does_not_misfire_on_a_genuinely_corrupt_non_zip_file():
    """A file that isn't a valid zip at all (e.g. truncated/corrupt) must
    fall through to the real parser's own "could not read" error, not this
    guard's "too large" message — they're different failures."""
    result = chat_files.extract_attachment(
        "corrupt.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        b"this is not a zip file at all")
    assert result["kind"] == "error"
    assert "too large" not in result["message"]


def test_extract_truncates_a_huge_text_file_instead_of_blowing_the_context():
    huge = ("x" * (chat_files.MAX_TEXT_CHARS + 5000)).encode()
    result = chat_files.extract_attachment("huge.txt", "text/plain", huge)
    assert result["kind"] == "text"
    assert result["truncated"] is True
    assert len(result["text"]) == chat_files.MAX_TEXT_CHARS


# --------------------------------------------------------------------- #
# chat_agent._to_lc_messages — pure, no DB/LLM
# --------------------------------------------------------------------- #
def test_to_lc_messages_plain_history_no_attachment():
    from langchain_core.messages import AIMessage, HumanMessage
    history = [
        {"role": "user", "content": "hi"},
        {"role": "assistant", "content": "hello"},
        {"role": "user", "content": "how are you"},
    ]
    out = chat_agent._to_lc_messages(history, None)
    assert len(out) == 3
    assert isinstance(out[0], HumanMessage) and out[0].content == "hi"
    assert isinstance(out[1], AIMessage) and out[1].content == "hello"
    assert isinstance(out[2], HumanMessage) and out[2].content == "how are you"


def test_to_lc_messages_folds_text_attachment_into_last_turn_only():
    history = [
        {"role": "user", "content": "earlier message, unrelated"},
        {"role": "user", "content": "what does this say?"},
    ]
    attachment = {"kind": "text", "filename": "notes.txt", "text": "the secret is 42", "truncated": False}
    out = chat_agent._to_lc_messages(history, attachment)
    assert isinstance(out[0].content, str)   # earlier turn untouched, still plain text
    last = out[-1]
    assert isinstance(last.content, list)     # last turn became multimodal blocks
    joined = " ".join(b.get("text", "") for b in last.content if b.get("type") == "text")
    assert "what does this say?" in joined
    assert "the secret is 42" in joined
    assert "notes.txt" in joined


def test_to_lc_messages_folds_image_attachment_as_image_block():
    history = [{"role": "user", "content": "look at this"}]
    attachment = {"kind": "image", "mime": "image/png", "b64": "AAAA"}
    out = chat_agent._to_lc_messages(history, attachment)
    last = out[-1]
    kinds = [b["type"] for b in last.content]
    assert "image_url" in kinds
    img_block = next(b for b in last.content if b["type"] == "image_url")
    assert img_block["image_url"]["url"].startswith("data:image/png;base64,")


def test_to_lc_messages_allows_a_file_with_no_typed_caption():
    """Dropping a file with nothing typed is a real turn, not empty."""
    history = [{"role": "user", "content": ""}]
    attachment = {"kind": "image", "mime": "image/png", "b64": "AAAA"}
    out = chat_agent._to_lc_messages(history, attachment)
    assert len(out) == 1
    texts = [b["text"] for b in out[0].content if b.get("type") == "text"]
    assert any(texts)   # a default prompt was filled in, not silently dropped


def test_to_lc_messages_reports_a_failed_attachment_as_text_not_silently_dropped():
    history = [{"role": "user", "content": "what's in the file"}]
    attachment = {"kind": "error", "message": "Could not read x.doc: unsupported"}
    out = chat_agent._to_lc_messages(history, attachment)
    joined = " ".join(b.get("text", "") for b in out[-1].content if b.get("type") == "text")
    assert "could not be read" in joined.lower()
    assert "unsupported" in joined


# --------------------------------------------------------------------- #
# Endpoints
# --------------------------------------------------------------------- #
async def test_suggestions_endpoint_returns_simplified_shape(client, admin_token):
    h = auth_headers(admin_token)
    r = await client.get("/api/v1/agentic-chat/suggestions", headers=h)
    assert r.status_code == 200, r.text
    body = r.json()
    assert "enabled" in body
    assert "suggestions" not in body and "prompt_book" not in body


async def test_chat_stream_without_api_key_degrades_gracefully(client, admin_token):
    h = auth_headers(admin_token)
    r = await client.post(
        "/api/v1/agentic-chat/stream", headers=h,
        data={"messages": json.dumps([{"role": "user", "content": "hello"}])},
    )
    assert r.status_code == 200, r.text
    # No key configured in the test env -> a graceful token + done(error),
    # never a crash.
    assert "no_api_key" in r.text
    assert "AI provider" in r.text


async def test_chat_stream_enforces_the_access_toggle_server_side_not_just_in_the_ui(client, admin_token):
    """The AI Settings "Ask AI access" toggle used to only be checked by the
    frontend's lock screen — a non-admin token calling /stream directly (curl,
    no UI at all) bypassed it completely. Must be enforced here too."""
    from tests.conftest import login_2fa

    h = auth_headers(admin_token)
    u = await client.post("/api/v1/admin/users", headers=h, json={
        "username": "chat-others-person", "password": "pw12345678", "role": "user", "auth_mode": "captcha",
    })
    assert u.status_code == 201, u.text
    other_token = await login_2fa(client, "chat-others-person", "pw12345678")
    oh = auth_headers(other_token)

    # Default (enabled_for_others=False) -> blocked even with a valid,
    # correctly-scoped "user"-role token.
    r = await client.post(
        "/api/v1/agentic-chat/stream", headers=oh,
        data={"messages": json.dumps([{"role": "user", "content": "hello"}])},
    )
    assert r.status_code == 403, r.text

    # Admin turns it on -> the SAME token now streams.
    put = await client.put("/api/v1/agentic-chat/access", headers=h, json={"enabled_for_others": True})
    assert put.status_code == 200, put.text
    try:
        r2 = await client.post(
            "/api/v1/agentic-chat/stream", headers=oh,
            data={"messages": json.dumps([{"role": "user", "content": "hello"}])},
        )
        assert r2.status_code == 200, r2.text
    finally:
        await client.put("/api/v1/agentic-chat/access", headers=h, json={"enabled_for_others": False})


async def test_chat_stream_accepts_a_file_attachment(client, admin_token):
    """Deliberate reversal of the old chat's design (it had NO upload surface
    at all) — a generic chat now DOES accept an attached file, the same as
    ChatGPT/Claude. This just proves the endpoint accepts and reads it
    (still degrades gracefully afterwards with no API key in the test env)."""
    h = auth_headers(admin_token)
    r = await client.post(
        "/api/v1/agentic-chat/stream", headers=h,
        data={"messages": json.dumps([{"role": "user", "content": "summarize this"}])},
        files={"file": ("notes.txt", b"the quarterly numbers look good", "text/plain")},
    )
    assert r.status_code == 200, r.text
    assert "no_api_key" in r.text
